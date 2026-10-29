import sys
from pathlib import Path
from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QHBoxLayout, QVBoxLayout,
    QLabel, QPushButton, QFileDialog, QFrame, QStackedWidget, QLineEdit, QSpinBox,
    QTableWidget, QTableWidgetItem, QMessageBox, QProgressBar, QHeaderView)
from openpyxl import load_workbook, Workbook

STYLE = """
QWidget {background:#f5f7fb;color:#18233a;font-family:'Segoe UI';font-size:13px}
QFrame#sidebar {background:#101d35;border:0}
QFrame#sidebar QLabel {background:transparent;color:#f8fafc}
QPushButton#nav {text-align:left;background:transparent;color:#aebbd0;padding:13px 16px;border:0;border-radius:9px}
QPushButton#nav:checked,QPushButton#nav:hover {background:#263b60;color:white}
QPushButton#primary {background:#2563eb;color:white;border:0;border-radius:9px;padding:12px 19px;font-weight:600}
QPushButton#primary:hover {background:#1d4ed8}
QPushButton#primary:disabled {background:#aab7ce}
QPushButton#secondary {background:white;color:#1e3a5f;border:1px solid #dbe3ef;border-radius:9px;padding:11px 16px}
QFrame#panel {background:white;border:1px solid #e3e9f3;border-radius:14px}
QFrame#panel QLabel {background:transparent}
QLineEdit,QSpinBox {background:white;border:1px solid #d5dfeb;border-radius:8px;padding:9px}
QTableWidget {background:white;gridline-color:#edf1f6;border:1px solid #e3e9f3;border-radius:8px}
QHeaderView::section {background:#eff4fb;color:#334155;padding:8px;border:0}
QProgressBar {background:#e5ecf6;border:0;border-radius:5px;text-align:center}
QProgressBar::chunk {background:#2563eb;border-radius:5px}
"""

class SplitWorker(QThread):
    progress = Signal(int)
    finished_ok = Signal(str)
    failed = Signal(str)
    def __init__(self, source, destination, count):
        super().__init__()
        self.source, self.destination, self.count = source, destination, count

    def run(self):
        try:
            source = load_workbook(self.source, read_only=True, data_only=False)
            try:
                sheet = source.active
                total = max(0, sheet.max_row - 1)
                if total == 0:
                    raise ValueError("Selected worksheet has no data rows.")
                destination = Path(self.destination)
                destination.mkdir(parents=True, exist_ok=True)
                headers = next(sheet.iter_rows(values_only=True))
                part, written = 0, 0
                output = None
                for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True)):
                    if index % self.count == 0:
                        if output:
                            output.save(str(destination / f"Part_{part:03d}.xlsx"))
                            output.close()
                        part += 1
                        output = Workbook(write_only=True)
                        output.create_sheet("Data").append(headers)
                    output.active.append(row)
                    written += 1
                    if index % 100 == 0:
                        self.progress.emit(min(99, int(written * 100 / total)))
                if output:
                    output.save(str(destination / f"Part_{part:03d}.xlsx"))
                    output.close()
                self.progress.emit(100)
                self.finished_ok.emit(f"Created {part} files in {destination}")
            finally:
                source.close()
        except Exception as exc:
            self.failed.emit(str(exc))

class DropLabel(QLabel):
    dropped = Signal(str)
    def __init__(self):
        super().__init__("Drop an .xlsx file here, or click Browse")
        self.setAlignment(Qt.AlignCenter)
        self.setAcceptDrops(True)
        self.setStyleSheet("border:2px dashed #a8bce1;border-radius:12px;padding:25px;color:#52719e;background:#f8fbff")
    def dragEnterEvent(self, event):
        urls = event.mimeData().urls()
        if urls and urls[0].toLocalFile().lower().endswith(".xlsx"):
            event.acceptProposedAction()
    def dropEvent(self, event):
        urls = event.mimeData().urls()
        if urls:
            self.dropped.emit(urls[0].toLocalFile())
            event.acceptProposedAction()

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Split Data Pro | Data Processing Studio")
        self.resize(1150, 750)
        self.source = None
        self.worker = None
        shell = QWidget()
        self.setCentralWidget(shell)
        outer = QHBoxLayout(shell)
        outer.setContentsMargins(0,0,0,0)
        outer.setSpacing(0)
        sidebar = QFrame(objectName="sidebar")
        sidebar.setFixedWidth(232)
        menu = QVBoxLayout(sidebar)
        menu.setContentsMargins(18,28,18,22)
        brand = QLabel("◈  Split Data Pro")
        brand.setStyleSheet("font-size:20px;font-weight:700")
        menu.addWidget(brand)
        subtitle = QLabel("DATA PROCESSING STUDIO")
        subtitle.setStyleSheet("font-size:10px;color:#93a7c8")
        menu.addWidget(subtitle)
        menu.addSpacing(35)
        self.pages = QStackedWidget()
        self.nav = []
        for i, name in enumerate(["Dashboard","Split Data","Combine","Data Cleanup","File Tools"]):
            button = QPushButton("   "+name,objectName="nav")
            button.setCheckable(True)
            button.clicked.connect(lambda _, n=i: self.show_page(n))
            menu.addWidget(button)
            self.nav.append(button)
            self.pages.addWidget(self.make_page(name, i))
        menu.addStretch()
        menu.addWidget(QLabel("v0.1.0  •  Phase 1"))
        outer.addWidget(sidebar)
        outer.addWidget(self.pages, 1)
        self.show_page(0)
        self.statusBar().showMessage("Ready • All processing stays on your computer")

    def panel(self):
        frame = QFrame(objectName="panel")
        layout = QVBoxLayout(frame)
        layout.setContentsMargins(22,22,22,22)
        layout.setSpacing(14)
        return frame,layout

    def make_page(self, name, index):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(34,32,34,30)
        layout.setSpacing(18)
        title = QLabel(name)
        title.setStyleSheet("font-size:27px;font-weight:700")
        layout.addWidget(title)
        desc = QLabel("Organize, transform and export Excel data with confidence.")
        desc.setStyleSheet("color:#64748b")
        layout.addWidget(desc)
        if index == 0:
            frame, body = self.panel()
            heading = QLabel("Welcome to Split Data Pro")
            heading.setStyleSheet("font-size:21px;font-weight:650")
            body.addWidget(heading)
            body.addWidget(QLabel("Your local workspace for Excel data. Choose a tool to get started."))
            start = QPushButton("Open Split Data",objectName="primary")
            start.clicked.connect(lambda:self.show_page(1))
            body.addWidget(start,alignment=Qt.AlignLeft)
            layout.addWidget(frame)
        elif index == 1:
            frame, body = self.panel()
            self.drop = DropLabel()
            self.drop.dropped.connect(self.open_source)
            body.addWidget(self.drop)
            browse = QPushButton("Browse Excel File",objectName="secondary")
            browse.clicked.connect(self.browse)
            body.addWidget(browse,alignment=Qt.AlignLeft)
            self.file_info = QLabel("No workbook selected")
            body.addWidget(self.file_info)
            body.addWidget(QLabel("Rows per output workbook"))
            self.rows = QSpinBox()
            self.rows.setRange(1,1000000)
            self.rows.setValue(500)
            body.addWidget(self.rows)
            self.run_button = QPushButton("Split into Excel Workbooks",objectName="primary")
            self.run_button.clicked.connect(self.start_split)
            self.run_button.setEnabled(False)
            body.addWidget(self.run_button,alignment=Qt.AlignLeft)
            self.progress = QProgressBar()
            self.progress.setValue(0)
            body.addWidget(self.progress)
            self.preview = QTableWidget(0,0)
            self.preview.setMinimumHeight(220)
            self.preview.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
            body.addWidget(QLabel("Workbook Preview (first 20 rows)"))
            body.addWidget(self.preview)
            layout.addWidget(frame)
        else:
            frame,body = self.panel()
            body.addWidget(QLabel(f"{name} — coming in the next phase"))
            body.addWidget(QLabel("This workspace is planned but processing actions are not yet available."))
            layout.addWidget(frame)
        layout.addStretch()
        return page

    def show_page(self, index):
        self.pages.setCurrentIndex(index)
        for i, button in enumerate(self.nav):
            button.setChecked(i==index)

    def browse(self):
        path,_ = QFileDialog.getOpenFileName(self,"Choose Excel workbook","","Excel Workbook (*.xlsx)")
        if path:
            self.open_source(path)

    def open_source(self,path):
        try:
            wb=load_workbook(path,read_only=True,data_only=True)
            try:
                ws=wb.active
                sample=list(ws.iter_rows(max_row=21,values_only=True))
                headers=sample[0] if sample else ()
                self.preview.setColumnCount(len(headers))
                self.preview.setHorizontalHeaderLabels([str(h) if h is not None else "" for h in headers])
                self.preview.setRowCount(max(0,len(sample)-1))
                for r,row in enumerate(sample[1:]):
                    for c,value in enumerate(row):
                        self.preview.setItem(r,c,QTableWidgetItem("" if value is None else str(value)))
                self.source=path
                self.file_info.setText(f"{Path(path).name}  •  {ws.max_row-1:,} data rows  •  {ws.max_column} columns  •  Sheet: {ws.title}")
                self.run_button.setEnabled(True)
                self.statusBar().showMessage("Workbook loaded")
            finally:
                wb.close()
        except Exception as exc:
            self.source=None
            self.run_button.setEnabled(False)
            QMessageBox.warning(self,"Unable to open workbook",str(exc))

    def start_split(self):
        if not self.source:
            return
        folder=QFileDialog.getExistingDirectory(self,"Choose output folder")
        if not folder:
            return
        if list(Path(folder).glob("Part_*.xlsx")):
            QMessageBox.warning(self,"Output folder not empty","Choose a folder without Part_*.xlsx files to prevent overwriting.")
            return
        self.run_button.setEnabled(False)
        self.progress.setValue(0)
        self.worker=SplitWorker(self.source,folder,self.rows.value())
        self.worker.progress.connect(self.progress.setValue)
        self.worker.finished_ok.connect(self.done)
        self.worker.failed.connect(self.error)
        self.worker.start()
        self.statusBar().showMessage("Processing...")

    def done(self,message):
        self.run_button.setEnabled(True)
        self.statusBar().showMessage(message)
        QMessageBox.information(self,"Split complete",message)

    def error(self,message):
        self.run_button.setEnabled(True)
        self.statusBar().showMessage("Processing failed")
        QMessageBox.critical(self,"Split failed",message)

if __name__=="__main__":
    app=QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setStyleSheet(STYLE)
    window=MainWindow()
    window.show()
    sys.exit(app.exec())
